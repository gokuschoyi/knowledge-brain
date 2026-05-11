from __future__ import annotations

import csv
import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import pymupdf
import requests
from bs4 import BeautifulSoup
from pypdf import PdfReader

from apps.documents.models import Document
from apps.documents.services.text_cleaning import clean_text

_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)", re.MULTILINE)


def _append_block(full_text: str, block: str, separator: str = "\n") -> tuple[str, int, int]:
    if not block:
        return full_text, len(full_text), len(full_text)
    if full_text:
        full_text += separator
    start = len(full_text)
    full_text += block
    return full_text, start, len(full_text)


def _extract_pdf_text_from_bytes(pdf_bytes: bytes, extraction_source: str) -> tuple[str, list[dict], list[dict]]:
    pdf = pymupdf.open(stream=pdf_bytes, filetype="pdf")

    full_text = ""
    page_boundaries: list[dict] = []
    word_infos: list[dict] = []
    reading_order = 0

    for page_index, page in enumerate(pdf, 1):
        page_width = page.rect.width
        words = page.get_text("words")

        if not words:
            page_boundaries.append({"page": page_index, "start_char": len(full_text), "end_char": len(full_text)})
            continue

        words = sorted(words, key=lambda w: (w[5], w[6], w[7]))
        page_start = len(full_text)
        page_text = ""
        prev_line_key: tuple[int, int] | None = None

        for x0, y0, x1, y1, word_text, block_no, line_no, word_no in words:
            curr_line_key = (int(block_no), int(line_no))
            if prev_line_key is not None:
                page_text += "\n" if curr_line_key != prev_line_key else " "

            word_start = page_start + len(page_text)
            page_text += word_text
            word_end = page_start + len(page_text)

            word_infos.append(
                {
                    "text": word_text,
                    "start": word_start,
                    "end": word_end,
                    "page": page_index,
                    "bbox": [round(x0, 1), round(y0, 1), round(x1, 1), round(y1, 1), page_index, round(page_width, 1)],
                    "block_index": int(block_no),
                    "line_index": int(line_no),
                    "reading_order": reading_order,
                    "extraction_source": extraction_source,
                }
            )
            reading_order += 1
            prev_line_key = curr_line_key

        page_end = page_start + len(page_text)
        page_boundaries.append({"page": page_index, "start_char": page_start, "end_char": page_end})
        full_text += page_text + "\n\n"

    pdf.close()
    return full_text.strip(), page_boundaries, word_infos


def _extract_pdf_text_with_pypdf(document: Document) -> str:
    with document.raw_file.open("rb") as handle:
        reader = PdfReader(handle)
        page_text: list[str] = []
        for page in reader.pages:
            extracted = page.extract_text() or ""
            extracted = extracted.strip()
            if extracted:
                page_text.append(extracted)
        return clean_text("\n\n".join(page_text))


def _run_ocrmypdf(document: Document) -> bytes:
    if shutil.which("ocrmypdf") is None:
        raise RuntimeError(
            "OCRmyPDF is not installed. Install 'ocrmypdf' and 'tesseract-ocr' to OCR scanned PDFs."
        )

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_dir_path = Path(temp_dir)
        input_path = temp_dir_path / "input.pdf"
        output_path = temp_dir_path / "output_ocr.pdf"

        with document.raw_file.open("rb") as source_handle:
            input_path.write_bytes(source_handle.read())

        command = ["ocrmypdf", "--force-ocr", str(input_path), str(output_path)]
        result = subprocess.run(command, capture_output=True, text=True)
        if result.returncode != 0:
            stderr = result.stderr.strip() or result.stdout.strip() or "Unknown OCRmyPDF error"
            raise RuntimeError(f"OCRmyPDF failed: {stderr}")
        return output_path.read_bytes()


def _extract_pdf_text(document: Document) -> tuple[str, list[dict], list[dict], str]:
    with document.raw_file.open("rb") as handle:
        pdf_bytes = handle.read()

    pymupdf_text, page_boundaries, word_infos = _extract_pdf_text_from_bytes(
        pdf_bytes,
        extraction_source="native_pdf",
    )
    if len(pymupdf_text) >= 80:
        return pymupdf_text, page_boundaries, word_infos, "native_pdf"

    pypdf_text = _extract_pdf_text_with_pypdf(document)
    if len(pypdf_text) >= 80:
        return pypdf_text, [], [], "native_pdf"

    ocr_pdf_bytes = _run_ocrmypdf(document)
    ocr_text, ocr_page_boundaries, ocr_word_infos = _extract_pdf_text_from_bytes(
        ocr_pdf_bytes,
        extraction_source="ocr",
    )
    return ocr_text, ocr_page_boundaries, ocr_word_infos, "ocr"


def _extract_docx_text(file_bytes: bytes) -> tuple[str, dict]:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(file_bytes))
    text = ""
    paragraphs: list[dict] = []
    headings: list[dict] = []
    paragraph_index = 0

    for para in doc.paragraphs:
        block = para.text.strip()
        if not block:
            continue
        text, start, end = _append_block(text, block, separator="\n\n")
        paragraphs.append({"index": paragraph_index, "start_char": start, "end_char": end})
        style_name = getattr(getattr(para, "style", None), "name", "") or ""
        if style_name.lower().startswith("heading"):
            level_match = re.search(r"(\d+)", style_name)
            headings.append(
                {
                    "heading": block,
                    "level": int(level_match.group(1)) if level_match else 1,
                    "start_char": start,
                }
            )
        paragraph_index += 1

    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows, 1):
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if not cells:
                continue
            block = " | ".join(cells)
            text, start, end = _append_block(text, block, separator="\n")
            paragraphs.append(
                {
                    "index": paragraph_index,
                    "start_char": start,
                    "end_char": end,
                    "table_index": table_index,
                    "row_index": row_index,
                }
            )
            paragraph_index += 1

    return text, {"paragraphs": paragraphs, "headings": headings}


def _extract_xlsx_text(file_bytes: bytes) -> tuple[str, dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    text = ""
    rows: list[dict] = []
    sheet_ranges: list[dict] = []

    for sheet in wb.worksheets:
        sheet_start = len(text)
        text, _, _ = _append_block(text, f"[Sheet: {sheet.title}]", separator="\n\n")
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
            cells = [str(cell) for cell in row if cell is not None and str(cell).strip()]
            if not cells:
                continue
            block = " | ".join(cells)
            text, start, end = _append_block(text, block, separator="\n")
            rows.append(
                {
                    "sheet_name": sheet.title,
                    "row_start": row_index,
                    "row_end": row_index,
                    "column_start": 1,
                    "column_end": len(cells),
                    "start_char": start,
                    "end_char": end,
                }
            )
        sheet_ranges.append({"sheet_name": sheet.title, "start_char": sheet_start, "end_char": len(text)})

    return text, {"sheet_rows": rows, "sheet_ranges": sheet_ranges}


def _extract_csv_text(file_bytes: bytes) -> tuple[str, dict]:
    text = file_bytes.decode("utf-8", errors="ignore")
    parsed_rows = list(csv.reader(io.StringIO(text)))

    built_text = ""
    row_ranges: list[dict] = []
    for row_index, row in enumerate(parsed_rows, 1):
        cells = [cell.strip() for cell in row if cell and cell.strip()]
        if not cells:
            continue
        block = " | ".join(cells)
        built_text, start, end = _append_block(built_text, block, separator="\n")
        row_ranges.append(
            {
                "sheet_name": "CSV",
                "row_start": row_index,
                "row_end": row_index,
                "column_start": 1,
                "column_end": len(cells),
                "start_char": start,
                "end_char": end,
            }
        )
    return built_text, {"sheet_rows": row_ranges}


def _extract_pptx_text(file_bytes: bytes) -> tuple[str, dict]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(file_bytes))
    text = ""
    headings: list[dict] = []
    slides: list[dict] = []
    slide_text_ranges: list[dict] = []

    for slide_index, slide in enumerate(prs.slides, 1):
        slide_start = len(text)
        marker = f"[Slide {slide_index}]"
        text, marker_start, _ = _append_block(text, marker, separator="\n\n")
        headings.append({"heading": f"Slide {slide_index}", "level": 1, "start_char": marker_start})
        for shape_index, shape in enumerate(slide.shapes, 1):
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                block = para.text.strip()
                if not block:
                    continue
                text, start, end = _append_block(text, block, separator="\n")
                slide_text_ranges.append(
                    {
                        "slide_number": slide_index,
                        "shape_index": shape_index,
                        "start_char": start,
                        "end_char": end,
                    }
                )
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                block = f"[Notes] {notes}"
                text, start, end = _append_block(text, block, separator="\n")
                slide_text_ranges.append(
                    {
                        "slide_number": slide_index,
                        "shape_index": None,
                        "start_char": start,
                        "end_char": end,
                        "is_notes": True,
                    }
                )
        slides.append({"slide_number": slide_index, "start_char": slide_start, "end_char": len(text)})

    return text, {"slides": slides, "slide_text_ranges": slide_text_ranges, "headings": headings}


def _detect_markdown_headings(text: str) -> list[dict]:
    headings = []
    for match in _MD_HEADING_RE.finditer(text):
        headings.append(
            {
                "heading": match.group(2).strip(),
                "level": len(match.group(1)),
                "start_char": match.start(),
            }
        )
    return headings


def extract_text(document: Document) -> dict:
    page_boundaries: list[dict] = []
    headings: list[dict] = []
    doc_words: list[dict] = []
    file_extension: str | None = None
    structure_metadata: dict = {}
    extraction_source = "native"

    if document.source_type == Document.SOURCE_TEXT:
        text = clean_text(document.raw_text)
        headings = _detect_markdown_headings(text)
        structure_metadata = {"headings": headings}

    elif document.source_type == Document.SOURCE_URL and document.url:
        response = requests.get(document.url, timeout=10)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, "html.parser")
        text = clean_text(soup.get_text("\n", strip=True))
        structure_metadata = {}

    elif document.source_type == Document.SOURCE_FILE and document.raw_file:
        suffix = Path(document.raw_file.name).suffix.lower()
        file_extension = suffix.lstrip(".")

        if suffix in {".txt", ".md"}:
            with document.raw_file.open("rb") as handle:
                with io.TextIOWrapper(handle, encoding="utf-8", errors="ignore") as text_handle:
                    text = clean_text(text_handle.read())
            headings = _detect_markdown_headings(text) if suffix == ".md" else []
            structure_metadata = {"headings": headings}
        elif suffix == ".pdf":
            text, page_boundaries, doc_words, extraction_source = _extract_pdf_text(document)
            structure_metadata = {
                "page_boundaries": page_boundaries,
                "headings": [],
                "extraction_source": extraction_source,
            }
        else:
            with document.raw_file.open("rb") as handle:
                file_bytes = handle.read()
            if suffix == ".docx":
                text, structure_metadata = _extract_docx_text(file_bytes)
                headings = structure_metadata.get("headings", [])
            elif suffix == ".xlsx":
                text, structure_metadata = _extract_xlsx_text(file_bytes)
            elif suffix == ".csv":
                text, structure_metadata = _extract_csv_text(file_bytes)
            elif suffix == ".pptx":
                text, structure_metadata = _extract_pptx_text(file_bytes)
                headings = structure_metadata.get("headings", [])
            else:
                text = ""
    else:
        text = ""

    return {
        "text": text,
        "page_boundaries": page_boundaries,
        "headings": headings,
        "file_extension": file_extension,
        "document_words": doc_words,
        "structure_metadata": structure_metadata,
        "extraction_source": extraction_source,
    }
